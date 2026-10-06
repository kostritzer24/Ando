"""RF-19/RF-20: generar y cargar la plantilla de calificaciones. Cada
celda se clasifica antes de guardar nada: "crear" (sin nota todavía),
"sin_cambio", "correccion" (otra nota, dentro del plazo de entrega: se
corrige directo, RN-05) o "modificacion" (otra nota, después del plazo:
RN-07, se trata como solicitud que autoriza Dirección)."""

import io
import zipfile

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from apps.students.models import Enrollment

from ..domain.grade_change import puede_corregirse_sin_autorizacion
from ..domain.template import (
    CeldaPlantillaInvalida,
    PlantillaNoCorresponde,
    validar_codigo,
    validar_encabezados,
    validar_identidad,
    validar_punteo,
)
from ..models import Activity, Grade, GradeChangeRequest
from .grade import corregir_nota_en_plazo, registrar_punteo
from .grade_change_request import solicitar_modificacion

ACCION_CREAR = "crear"
ACCION_MODIFICACION = "modificacion"
ACCION_CORRECCION = "correccion"
ACCION_SIN_CAMBIO = "sin_cambio"

# Hoja oculta con los identificadores de lo que se descargó: asignación,
# unidad y, en orden de columna, cada actividad. Es lo que permite
# rechazar una plantilla de otro curso, de otra unidad o desactualizada.
HOJA_IDENTIDAD = "_plantilla"


class UnidadSinActividades(Exception):
    pass


class ArchivoIlegible(Exception):
    pass


def _actividades_de(assignment, unit) -> list[Activity]:
    actividades = list(
        Activity.objects.filter(assignment=assignment, unit=unit, is_active=True).order_by(
            "due_date", "id"
        )
    )
    if not actividades:
        raise UnidadSinActividades(
            "Esta unidad todavía no tiene actividades definidas para esta asignación."
        )
    return actividades


def _encabezados(actividades: list[Activity]) -> list[str]:
    return ["Código", "Nombre"] + [f"{a.name} (máx {a.max_score})" for a in actividades]


def _inscripciones_de(assignment):
    return Enrollment.objects.filter(
        section=assignment.section, cycle=assignment.cycle, is_active=True
    ).select_related("student")


def generar_plantilla(*, assignment, unit) -> bytes:
    actividades = _actividades_de(assignment, unit)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Calificaciones"
    encabezados = _encabezados(actividades)
    hoja.append(encabezados)

    for inscripcion in _inscripciones_de(assignment).order_by("student__internal_code"):
        fila = [inscripcion.student.internal_code, inscripcion.student.nombre_completo()]
        fila += ["" for _ in actividades]
        hoja.append(fila)

    for columna in range(1, len(encabezados) + 1):
        hoja.column_dimensions[get_column_letter(columna)].width = 24

    # Código, nombre y encabezados quedan bloqueados (HU-19: no se
    # modifican); solo las celdas de punteo se pueden escribir. Sin
    # contraseña: es una guía para quien llena el archivo, la validación
    # real es la del servidor al cargarlo.
    # El código como texto: un código numérico perdería sus ceros a la
    # izquierda si Excel lo convierte en número.
    for (celda,) in hoja.iter_rows(min_row=2, max_col=1):
        celda.number_format = "@"
    for fila in hoja.iter_rows(min_row=2, min_col=3, max_col=len(encabezados)):
        for celda in fila:
            celda.protection = Protection(locked=False)
    hoja.protection.sheet = True

    identidad = libro.create_sheet(HOJA_IDENTIDAD)
    identidad.append([str(assignment.public_id)])
    identidad.append([str(unit.public_id)])
    for actividad in actividades:
        identidad.append([str(actividad.public_id)])
    identidad.sheet_state = "veryHidden"

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _leer_identidad(libro) -> tuple[str | None, str | None, list[str]]:
    if HOJA_IDENTIDAD not in libro.sheetnames:
        return None, None, []
    valores = [
        str(fila[0]) for fila in libro[HOJA_IDENTIDAD].iter_rows(values_only=True) if fila[0]
    ]
    if len(valores) < 2:
        return None, None, []
    return valores[0], valores[1], valores[2:]


def validar_y_clasificar_plantilla(*, assignment, unit, archivo):
    """Devuelve `(filas_validas, errores)`. `filas_validas` es una lista
    de `(enrollment, [{"activity", "punteo", "accion"}, ...])`. Nunca
    guarda nada — sirve tanto para la vista previa como para la carga."""
    actividades = _actividades_de(assignment, unit)

    try:
        libro = load_workbook(archivo, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ArchivoIlegible(
            "No se pudo leer el archivo. "
            "Asegurate de que sea la plantilla en formato Excel (.xlsx)."
        ) from exc

    asignacion_recibida, unidad_recibida, actividades_recibidas = _leer_identidad(libro)
    hoja = libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    encabezados_recibidos = list(next(filas, ()))
    try:
        validar_identidad(
            asignacion_esperada=str(assignment.public_id),
            unidad_esperada=str(unit.public_id),
            actividades_esperadas=[str(a.public_id) for a in actividades],
            asignacion_recibida=asignacion_recibida,
            unidad_recibida=unidad_recibida,
            actividades_recibidas=actividades_recibidas,
        )
        validar_encabezados(esperados=_encabezados(actividades), recibidos=encabezados_recibidos)
    except PlantillaNoCorresponde as exc:
        return [], [str(exc)]

    inscripciones = {i.student.internal_code: i for i in _inscripciones_de(assignment)}
    # Una consulta para todas las notas y solicitudes pendientes de la
    # unidad, en vez de una por celda.
    notas_existentes = {
        (g.enrollment_id, g.activity_id): g
        for g in Grade.objects.filter(activity__in=actividades, is_active=True)
    }
    notas_con_pendiente = set(
        GradeChangeRequest.objects.filter(
            grade__activity__in=actividades,
            status=GradeChangeRequest.ESTADO_PENDIENTE,
            is_active=True,
        ).values_list("grade_id", flat=True)
    )

    en_plazo = puede_corregirse_sin_autorizacion(
        hoy=timezone.localdate(), fecha_entrega_notas=unit.grades_due_date
    )
    errores = []
    filas_validas = []
    filas_por_codigo: dict[str, int] = {}

    for numero, fila in enumerate(filas, start=2):
        if not fila or fila[0] in (None, ""):
            continue
        codigo = str(fila[0]).strip()

        try:
            validar_codigo(
                numero_fila=numero, internal_code=codigo, codigos_esperados=set(inscripciones)
            )
        except CeldaPlantillaInvalida as exc:
            errores.append(str(exc))
            continue

        if codigo in filas_por_codigo:
            errores.append(
                f"Fila {numero}: el código '{codigo}' ya aparece en la fila "
                f"{filas_por_codigo[codigo]}. Cada estudiante va una sola vez."
            )
            continue
        filas_por_codigo[codigo] = numero

        inscripcion = inscripciones[codigo]
        celdas = []
        fila_tiene_error = False

        for indice, actividad in enumerate(actividades):
            columna = 2 + indice
            valor = fila[columna] if len(fila) > columna else None
            if valor in (None, ""):
                continue

            try:
                punteo = validar_punteo(
                    numero_fila=numero,
                    nombre_actividad=actividad.name,
                    valor=valor,
                    max_score=actividad.max_score,
                )
            except CeldaPlantillaInvalida as exc:
                errores.append(str(exc))
                fila_tiene_error = True
                continue

            grade_existente = notas_existentes.get((inscripcion.id, actividad.id))
            if grade_existente is None:
                accion = ACCION_CREAR
            elif grade_existente.current_score == punteo:
                accion = ACCION_SIN_CAMBIO
            elif grade_existente.id in notas_con_pendiente:
                errores.append(
                    f"Fila {numero}: '{actividad.name}' ya tiene una corrección pendiente "
                    "de que Dirección la resuelva. Dejá esa celda como estaba."
                )
                fila_tiene_error = True
                continue
            else:
                accion = ACCION_CORRECCION if en_plazo else ACCION_MODIFICACION

            celdas.append(
                {
                    "activity": actividad,
                    "punteo": punteo,
                    "accion": accion,
                    "grade": grade_existente,
                }
            )

        if not fila_tiene_error:
            filas_validas.append((inscripcion, celdas))

    if errores:
        return [], errores
    return filas_validas, []


@transaction.atomic
def aplicar_plantilla(*, filas_validas, recorded_by) -> dict:
    """Solo se llama después de que `validar_y_clasificar_plantilla` no
    encontró errores. Todo o nada: si una celda falla (por ejemplo, otra
    persona registró esa nota entre la vista previa y la carga), no queda
    guardada ninguna."""
    creados = 0
    correcciones = 0
    solicitudes = 0
    for inscripcion, celdas in filas_validas:
        for celda in celdas:
            if celda["accion"] == ACCION_CREAR:
                registrar_punteo(
                    enrollment=inscripcion,
                    activity=celda["activity"],
                    raw_score=celda["punteo"],
                    recorded_by=recorded_by,
                    source=Grade.ORIGEN_PLANTILLA,
                )
                creados += 1
            elif celda["accion"] == ACCION_CORRECCION:
                corregir_nota_en_plazo(
                    celda["grade"], nuevo_punteo=celda["punteo"], usuario=recorded_by
                )
                correcciones += 1
            elif celda["accion"] == ACCION_MODIFICACION:
                solicitar_modificacion(
                    grade=celda["grade"],
                    requested_score=celda["punteo"],
                    reason="Cargado de nuevo por plantilla (RN-07).",
                    requested_by=recorded_by,
                )
                solicitudes += 1
    return {
        "creados": creados,
        "correcciones": correcciones,
        "solicitudes_de_modificacion": solicitudes,
    }
