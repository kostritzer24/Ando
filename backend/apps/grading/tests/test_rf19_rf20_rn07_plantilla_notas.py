import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.tests.factories import RoleFactory, UserFactory
from apps.catalog.tests.factories import GradingUnitFactory
from apps.grading.models import Grade, GradeChangeRequest
from apps.grading.services.template import generar_plantilla
from apps.students.tests.factories import EnrollmentFactory, StudentFactory

from .factories import ActivityFactory, ActivityTypeFactory, TeacherAssignmentFactory


def _docente_con_actividades_y_estudiantes(cantidad=2):
    rol = RoleFactory(name="Docente", permissions={"notas": "editar"})
    docente = UserFactory(role=rol)
    asignacion = TeacherAssignmentFactory(teacher=docente)
    unidad = GradingUnitFactory(cycle=asignacion.cycle)
    tipo = ActivityTypeFactory(name="Prueba corta")
    actividad = ActivityFactory(
        assignment=asignacion, unit=unidad, activity_type=tipo, name="Prueba corta 1", max_score=10
    )
    inscripciones = []
    for indice in range(cantidad):
        estudiante = StudentFactory(internal_code=f"ES{indice + 1:03d}")
        inscripciones.append(
            EnrollmentFactory(
                student=estudiante, section=asignacion.section, cycle=asignacion.cycle
            )
        )
    return docente, asignacion, unidad, actividad, inscripciones


def _archivo_xlsx(asignacion, unidad, filas: list[tuple]) -> SimpleUploadedFile:
    """Parte de la plantilla real que descarga el docente (con su hoja de
    identidad) y la llena con `filas`, como lo haría en Excel."""
    libro = load_workbook(io.BytesIO(generar_plantilla(assignment=asignacion, unit=unidad)))
    hoja = libro["Calificaciones"]
    hoja.delete_rows(2, hoja.max_row)
    for fila in filas:
        hoja.append(list(fila))
    return _como_subida(libro)


def _como_subida(libro) -> SimpleUploadedFile:
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        "notas.xlsx",
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.mark.django_db
def test_rf19_generar_plantilla_incluye_actividad_y_estudiantes():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(2)
    )

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.get(f"/api/v1/grades/template/{asignacion.public_id}/{unidad.public_id}/")

    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.content))
    hoja = libro.active
    encabezados = [c.value for c in next(hoja.iter_rows(max_row=1))]
    assert "Prueba corta 1" in encabezados[2]
    codigos = [fila[0].value for fila in hoja.iter_rows(min_row=2)]
    assert codigos == ["ES001", "ES002"]


@pytest.mark.django_db
def test_rf20_preview_clasifica_las_celdas_sin_guardar_nada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = _archivo_xlsx(asignacion, unidad, [("ES001", "Estudiante Uno", "8")])

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/preview/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 200
    assert respuesta.data["resumen"]["crear"] == 1
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_rf20_cargar_plantilla_valida_crea_las_notas():
    docente, asignacion, unidad, actividad, inscripciones = _docente_con_actividades_y_estudiantes(
        2
    )
    archivo = _archivo_xlsx(
        asignacion, unidad, [("ES001", "Estudiante Uno", "8"), ("ES002", "Estudiante Dos", "10")]
    )

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 201
    assert respuesta.data["creados"] == 2
    assert Grade.objects.filter(activity=actividad, source="plantilla").count() == 2


@pytest.mark.django_db
def test_rn07_recargar_la_plantilla_con_otra_nota_crea_una_solicitud_de_modificacion():
    docente, asignacion, unidad, actividad, inscripciones = _docente_con_actividades_y_estudiantes(
        1
    )
    Grade.objects.create(
        enrollment=inscripciones[0],
        activity=actividad,
        raw_score=7,
        current_score=7,
        recorded_by=docente,
    )
    archivo = _archivo_xlsx(asignacion, unidad, [("ES001", "Estudiante Uno", "9")])

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 201
    assert respuesta.data["solicitudes_de_modificacion"] == 1
    assert respuesta.data["creados"] == 0

    calificacion = Grade.objects.get(activity=actividad, enrollment=inscripciones[0])
    assert calificacion.current_score == 7  # sin cambios hasta que Dirección apruebe
    assert calificacion.raw_score == 7
    assert GradeChangeRequest.objects.filter(grade=calificacion, requested_score=9).exists()


@pytest.mark.django_db
def test_rn07_recargar_con_la_misma_nota_no_genera_nada():
    docente, asignacion, unidad, actividad, inscripciones = _docente_con_actividades_y_estudiantes(
        1
    )
    Grade.objects.create(
        enrollment=inscripciones[0],
        activity=actividad,
        raw_score=7,
        current_score=7,
        recorded_by=docente,
    )
    archivo = _archivo_xlsx(asignacion, unidad, [("ES001", "Estudiante Uno", "7")])

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 201
    assert respuesta.data == {"creados": 0, "correcciones": 0, "solicitudes_de_modificacion": 0}
    assert GradeChangeRequest.objects.count() == 0


@pytest.mark.django_db
def test_rf20_codigo_ajeno_no_guarda_nada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = _archivo_xlsx(asignacion, unidad, [("ES999", "No pertenece", "8")])

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 400
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_rechaza_archivos_que_no_son_xlsx():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo_falso = SimpleUploadedFile(
        "notas.xlsm", b"contenido", content_type="application/octet-stream"
    )

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {
            "assignment": str(asignacion.public_id),
            "unit": str(unidad.public_id),
            "file": archivo_falso,
        },
        format="multipart",
    )

    assert respuesta.status_code == 400


@pytest.mark.django_db
def test_docente_no_puede_cargar_plantilla_de_una_asignacion_ajena_acceso_no_autorizado():
    rol = RoleFactory(name="Docente", permissions={"notas": "editar"})
    docente = UserFactory(role=rol)
    _otro, asignacion_ajena, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = _archivo_xlsx(asignacion_ajena, unidad, [("ES001", "Estudiante Uno", "8")])

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.post(
        "/api/v1/grades/template/upload/",
        {
            "assignment": str(asignacion_ajena.public_id),
            "unit": str(unidad.public_id),
            "file": archivo,
        },
        format="multipart",
    )

    assert respuesta.status_code == 403


def _cargar(docente, asignacion, unidad, archivo, ruta="upload"):
    client = APIClient()
    client.force_authenticate(user=docente)
    return client.post(
        f"/api/v1/grades/template/{ruta}/",
        {"assignment": str(asignacion.public_id), "unit": str(unidad.public_id), "file": archivo},
        format="multipart",
    )


@pytest.mark.django_db
def test_hu19_la_plantilla_bloquea_codigo_y_nombre_y_deja_escribir_los_punteos():
    _docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    libro = load_workbook(io.BytesIO(generar_plantilla(assignment=asignacion, unit=unidad)))
    hoja = libro["Calificaciones"]

    assert hoja.protection.sheet is True
    assert hoja["A2"].protection.locked is True
    assert hoja["C2"].protection.locked is False
    assert libro["_plantilla"].sheet_state == "veryHidden"


@pytest.mark.django_db
def test_hu20_un_archivo_armado_a_mano_sin_la_plantilla_no_guarda_nada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    libro = openpyxl.Workbook()
    libro.active.append(["Código", "Nombre", "Prueba corta 1 (máx 10.00)"])
    libro.active.append(["ES001", "Estudiante Uno", 8])

    respuesta = _cargar(docente, asignacion, unidad, _como_subida(libro))

    assert respuesta.status_code == 400
    assert "no es una plantilla descargada" in respuesta.data["errores"][0]
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_hu20_columnas_cambiadas_de_lugar_no_guardan_en_la_actividad_equivocada():
    docente, asignacion, unidad, actividad, _inscripciones = _docente_con_actividades_y_estudiantes(
        1
    )
    ActivityFactory(
        assignment=asignacion, unit=unidad, name="Examen", max_score=50, due_date="2026-02-20"
    )
    libro = load_workbook(io.BytesIO(generar_plantilla(assignment=asignacion, unit=unidad)))
    hoja = libro["Calificaciones"]
    hoja["C1"], hoja["D1"] = hoja["D1"].value, hoja["C1"].value  # el docente invierte columnas
    hoja["C2"], hoja["D2"] = 40, 9

    respuesta = _cargar(docente, asignacion, unidad, _como_subida(libro))

    assert respuesta.status_code == 400
    assert "columnas" in respuesta.data["errores"][0]
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_hu20_plantilla_de_otra_unidad_es_rechazada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    otra_unidad = GradingUnitFactory(cycle=asignacion.cycle, number=unidad.number + 1)
    ActivityFactory(assignment=asignacion, unit=otra_unidad, name="Prueba corta 1", max_score=10)
    archivo = _archivo_xlsx(asignacion, otra_unidad, [("ES001", "Estudiante Uno", "8")])

    respuesta = _cargar(docente, asignacion, unidad, archivo)

    assert respuesta.status_code == 400
    assert "otra unidad" in respuesta.data["errores"][0]
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_hu20_plantilla_desactualizada_tras_agregar_una_actividad_es_rechazada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = _archivo_xlsx(asignacion, unidad, [("ES001", "Estudiante Uno", "8")])
    ActivityFactory(assignment=asignacion, unit=unidad, name="Examen", max_score=50)

    respuesta = _cargar(docente, asignacion, unidad, archivo)

    assert respuesta.status_code == 400
    assert "cambiaron" in respuesta.data["errores"][0]


@pytest.mark.django_db
def test_rf20_codigo_repetido_es_un_error_de_fila_y_no_guarda_nada():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = _archivo_xlsx(
        asignacion, unidad, [("ES001", "Estudiante Uno", "8"), ("ES001", "Estudiante Uno", "3")]
    )

    respuesta = _cargar(docente, asignacion, unidad, archivo)

    assert respuesta.status_code == 400
    assert "Fila 3" in respuesta.data["errores"][0]
    assert not Grade.objects.exists()


@pytest.mark.django_db
def test_rf20_un_archivo_que_no_es_excel_responde_400_no_500():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = SimpleUploadedFile("notas.xlsx", b"esto no es un zip", content_type="text/plain")

    respuesta = _cargar(docente, asignacion, unidad, archivo, ruta="preview")

    assert respuesta.status_code == 400


@pytest.mark.django_db
def test_rf20_rechaza_un_archivo_demasiado_pesado():
    docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    archivo = SimpleUploadedFile("notas.xlsx", b"0" * (2 * 1024 * 1024 + 1))

    respuesta = _cargar(docente, asignacion, unidad, archivo)

    assert respuesta.status_code == 400
    assert "pesa demasiado" in str(respuesta.data)


@pytest.mark.django_db
def test_rn07_una_celda_con_correccion_pendiente_no_crea_otra_solicitud():
    docente, asignacion, unidad, actividad, inscripciones = _docente_con_actividades_y_estudiantes(
        1
    )
    _cargar(docente, asignacion, unidad, _archivo_xlsx(asignacion, unidad, [("ES001", "Uno", "6")]))
    _cargar(docente, asignacion, unidad, _archivo_xlsx(asignacion, unidad, [("ES001", "Uno", "8")]))

    respuesta = _cargar(
        docente, asignacion, unidad, _archivo_xlsx(asignacion, unidad, [("ES001", "Uno", "9")])
    )

    assert respuesta.status_code == 400
    assert "pendiente" in respuesta.data["errores"][0]
    assert GradeChangeRequest.objects.count() == 1


@pytest.mark.django_db
def test_rn05_rn07_dentro_del_plazo_recargar_la_plantilla_corrige_directo():
    docente, asignacion, unidad, actividad, _inscripciones = _docente_con_actividades_y_estudiantes(
        1
    )
    unidad.grades_due_date = "2099-01-01"
    unidad.save()
    _cargar(docente, asignacion, unidad, _archivo_xlsx(asignacion, unidad, [("ES001", "Uno", "6")]))

    respuesta = _cargar(
        docente, asignacion, unidad, _archivo_xlsx(asignacion, unidad, [("ES001", "Uno", "8")])
    )

    assert respuesta.status_code == 201
    assert respuesta.data["correcciones"] == 1
    assert GradeChangeRequest.objects.count() == 0
    assert Grade.objects.get(activity=actividad).current_score == 8


@pytest.mark.django_db
def test_rf19_el_codigo_queda_como_texto_para_no_perder_ceros():
    _docente, asignacion, unidad, _actividad, _inscripciones = (
        _docente_con_actividades_y_estudiantes(1)
    )
    libro = load_workbook(io.BytesIO(generar_plantilla(assignment=asignacion, unit=unidad)))

    assert libro["Calificaciones"]["A2"].number_format == "@"
