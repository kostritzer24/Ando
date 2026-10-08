import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.tests.factories import RoleFactory, UserFactory
from apps.attendance.models import Attendance
from apps.catalog.tests.factories import SchoolCycleFactory, SectionFactory
from apps.students.tests.factories import EnrollmentFactory, StudentFactory


def _direccion():
    rol = RoleFactory(name="Dirección", permissions={"asistencia": "editar"})
    return UserFactory(role=rol)


def _seccion_taller_con_estudiantes(cantidad=2):
    ciclo = SchoolCycleFactory()
    seccion = SectionFactory(cycle=ciclo, type="taller", grade="Taller de panadería")
    inscripciones = []
    for indice in range(cantidad):
        estudiante = StudentFactory(internal_code=f"ES{indice + 1:03d}")
        inscripciones.append(EnrollmentFactory(student=estudiante, section=seccion, cycle=ciclo))
    return seccion, inscripciones


def _construir_archivo_xlsx(filas: list[tuple]) -> SimpleUploadedFile:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Código", "Nombre", "Estado"])
    for fila in filas:
        hoja.append(list(fila))
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        "asistencia.xlsx",
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.mark.django_db
def test_rf21_generar_plantilla_incluye_a_los_inscritos():
    direccion = _direccion()
    seccion, inscripciones = _seccion_taller_con_estudiantes(2)

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.get(f"/api/v1/attendance/template/{seccion.public_id}/2026-02-03/")

    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.content))
    hoja = libro.active
    codigos = [fila[0].value for fila in hoja.iter_rows(min_row=2) if fila[0].value]
    assert codigos == ["ES001", "ES002"]


@pytest.mark.django_db
def test_rf21_no_genera_plantilla_para_una_seccion_academica():
    direccion = _direccion()
    ciclo = SchoolCycleFactory()
    seccion = SectionFactory(cycle=ciclo, type="academica")

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.get(f"/api/v1/attendance/template/{seccion.public_id}/2026-02-03/")

    assert respuesta.status_code == 400


@pytest.mark.django_db
def test_rf21_cargar_plantilla_valida_crea_la_asistencia():
    direccion = _direccion()
    seccion, inscripciones = _seccion_taller_con_estudiantes(2)
    archivo = _construir_archivo_xlsx(
        [("ES001", "Estudiante Uno", "presente"), ("ES002", "Estudiante Dos", "ausente")]
    )

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.post(
        "/api/v1/attendance/template/upload/",
        {"section": str(seccion.public_id), "date": "2026-02-03", "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 201
    assert respuesta.data["creados"] == 2
    assert Attendance.objects.filter(enrollment__in=inscripciones, source="plantilla").count() == 2


@pytest.mark.django_db
def test_rf21_plantilla_con_codigo_ajeno_no_guarda_nada():
    direccion = _direccion()
    seccion, inscripciones = _seccion_taller_con_estudiantes(1)
    archivo = _construir_archivo_xlsx(
        [("ES001", "Estudiante Uno", "presente"), ("ES999", "No pertenece", "presente")]
    )

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.post(
        "/api/v1/attendance/template/upload/",
        {"section": str(seccion.public_id), "date": "2026-02-03", "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 400
    assert "errores" in respuesta.data
    assert not Attendance.objects.exists()


@pytest.mark.django_db
def test_rf21_no_se_puede_registrar_dos_veces_el_mismo_dia_por_plantilla():
    direccion = _direccion()
    seccion, inscripciones = _seccion_taller_con_estudiantes(1)
    Attendance.objects.create(
        enrollment=inscripciones[0],
        date="2026-02-03",
        status=Attendance.ESTADO_PRESENTE,
        recorded_by=direccion,
    )
    archivo = _construir_archivo_xlsx([("ES001", "Estudiante Uno", "ausente")])

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.post(
        "/api/v1/attendance/template/upload/",
        {"section": str(seccion.public_id), "date": "2026-02-03", "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 400
    assert Attendance.objects.count() == 1


@pytest.mark.django_db
def test_rf21_rechaza_archivos_que_no_son_xlsx():
    direccion = _direccion()
    seccion, _inscripciones = _seccion_taller_con_estudiantes(1)
    archivo_falso = SimpleUploadedFile(
        "asistencia.xlsm", b"contenido cualquiera", content_type="application/octet-stream"
    )

    client = APIClient()
    client.force_authenticate(user=direccion)

    respuesta = client.post(
        "/api/v1/attendance/template/upload/",
        {"section": str(seccion.public_id), "date": "2026-02-03", "file": archivo_falso},
        format="multipart",
    )

    assert respuesta.status_code == 400
    assert not Attendance.objects.exists()


@pytest.mark.django_db
def test_tallerista_no_puede_cargar_plantilla_de_una_seccion_ajena_acceso_no_autorizado():
    rol_tallerista = RoleFactory(name="Tallerista", permissions={"asistencia": "editar"})
    tallerista = UserFactory(role=rol_tallerista)
    seccion, _inscripciones = _seccion_taller_con_estudiantes(1)
    archivo = _construir_archivo_xlsx([("ES001", "Estudiante Uno", "presente")])

    client = APIClient()
    client.force_authenticate(user=tallerista)

    respuesta = client.post(
        "/api/v1/attendance/template/upload/",
        {"section": str(seccion.public_id), "date": "2026-02-03", "file": archivo},
        format="multipart",
    )

    assert respuesta.status_code == 403


@pytest.mark.django_db
def test_rnf04_un_docente_sin_asignacion_no_descarga_la_nomina_de_un_taller():
    """Hallazgo B-001: la descarga no revisaba la asignación (la subida sí)."""
    rol = RoleFactory(name="Docente", permissions={"asistencia": "ver"})
    docente = UserFactory(role=rol)
    seccion, _ = _seccion_taller_con_estudiantes(2)

    client = APIClient()
    client.force_authenticate(user=docente)

    respuesta = client.get(f"/api/v1/attendance/template/{seccion.public_id}/2026-02-03/")

    assert respuesta.status_code == 403
