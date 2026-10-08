"""RF-21: generar y cargar la plantilla de asistencia de talleres.
Biblioteca confirmada en ADR-0005 (openpyxl). Se lee con `data_only=True`
para no ejecutar fórmulas (sección 14.4 del prompt maestro); el llamador
(la vista) rechaza cualquier archivo que no sea `.xlsx` antes de llegar
acá, para no procesar macros de un `.xlsm`."""

import io
from datetime import date

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from apps.catalog.models import Section
from apps.students.models import Enrollment

from ..domain.template import FilaPlantillaInvalida, validar_fila
from ..models import Attendance
from .attendance import registrar_asistencia

_ENCABEZADOS = ["Código", "Nombre", "Estado (presente/ausente/justificado)"]


class SeccionNoEsDeTaller(Exception):
    pass


def _validar_seccion_de_taller(section: Section) -> None:
    if section.type != Section.TIPO_TALLER:
        raise SeccionNoEsDeTaller("La plantilla de asistencia es solo para secciones de taller.")


def generar_plantilla(*, section: Section) -> bytes:
    _validar_seccion_de_taller(section)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Asistencia"
    hoja.append(_ENCABEZADOS)

    inscripciones = (
        Enrollment.objects.filter(section=section, is_active=True)
        .select_related("student")
        .order_by("student__internal_code")
    )
    for inscripcion in inscripciones:
        # Código y nombre son de solo lectura en la plantilla real (no se
        # pueden modificar, HU-19/20) — acá solo se generan como guía.
        hoja.append([inscripcion.student.internal_code, inscripcion.student.nombre_completo(), ""])

    for columna in range(1, len(_ENCABEZADOS) + 1):
        hoja.column_dimensions[get_column_letter(columna)].width = 30

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


@transaction.atomic
def procesar_plantilla(
    *, section: Section, fecha: date, archivo, recorded_by
) -> tuple[list[Attendance], list[str]]:
    """Valida todas las filas primero; si hay algún error, no guarda nada
    y devuelve la lista de errores (la "vista previa" de la sección 14.4).
    Si no hay errores, crea la asistencia de cada fila."""
    _validar_seccion_de_taller(section)

    libro = load_workbook(archivo, data_only=True, read_only=True)
    hoja = libro.active

    inscripciones = {
        inscripcion.student.internal_code: inscripcion
        for inscripcion in Enrollment.objects.filter(
            section=section, is_active=True
        ).select_related("student")
    }
    ya_registrados = set(
        Attendance.objects.filter(enrollment__section=section, date=fecha).values_list(
            "enrollment__student__internal_code", flat=True
        )
    )

    filas_validas = []
    errores = []
    for numero, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if not fila or fila[0] in (None, ""):
            continue
        codigo = str(fila[0]).strip()
        estado = str(fila[2]).strip().lower() if len(fila) > 2 and fila[2] else ""

        try:
            validar_fila(
                numero_fila=numero,
                internal_code=codigo,
                status=estado,
                codigos_esperados=set(inscripciones),
            )
        except FilaPlantillaInvalida as exc:
            errores.append(str(exc))
            continue

        if codigo in ya_registrados:
            errores.append(
                f"Fila {numero}: ya existe asistencia registrada para '{codigo}' ese día."
            )
            continue

        filas_validas.append((inscripciones[codigo], estado))

    if errores:
        return [], errores

    creados = [
        registrar_asistencia(
            enrollment=inscripcion,
            fecha=fecha,
            recorded_by=recorded_by,
            status=estado,
            source=Attendance.ORIGEN_PLANTILLA,
        )
        for inscripcion, estado in filas_validas
    ]
    return creados, []
